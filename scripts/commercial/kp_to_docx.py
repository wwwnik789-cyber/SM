#!/usr/bin/env python3
"""Сборка клиентского КП из markdown в .docx (таблицы, заголовки, списки)."""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ACCENT = RGBColor(0x97, 0x1A, 0x81)
GREY = RGBColor(0x55, 0x55, 0x55)
MONO_TABLE = 8.5


def set_font(run, name="Manrope", size=10.5, bold=False, color=None, italic=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def add_rich(par, text, size=10.5, color=None, italic=False):
    """**жирный**, *курсив* и `код` внутри строки."""
    for chunk in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            r = par.add_run(chunk[2:-2])
            set_font(r, size=size, bold=True, color=color)
        elif chunk.startswith("*") and chunk.endswith("*"):
            r = par.add_run(chunk[1:-1])
            set_font(r, size=size, color=color, italic=True)
        elif chunk.startswith("`") and chunk.endswith("`"):
            r = par.add_run(chunk[1:-1])
            set_font(r, name="Consolas", size=size - 0.5, color=color)
        else:
            r = par.add_run(chunk)
            set_font(r, size=size, color=color, italic=italic)


def shade(cell, hexcolor):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.makeelement(qn("w:shd"), {})
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hexcolor)
    tc_pr.append(shd)


def build(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.6)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)

    style = doc.styles["Normal"]
    style.font.name = "Manrope"
    style.font.size = Pt(10.5)

    i = 0
    in_table = False
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if not rows:
                continue
            ncols = max(len(r) for r in rows)
            table = doc.add_table(rows=0, cols=ncols)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            # ширины колонок: № и служебные — узкие, текстовая — широкая
            if ncols == 4:
                widths = [Cm(0.9), Cm(11.6), Cm(2.2), Cm(2.5)]
            elif ncols == 3:
                widths = [Cm(8.2), Cm(6.0), Cm(3.0)]
            else:
                widths = [Cm(17.2 / ncols)] * ncols
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                for ci in range(ncols):
                    text = row[ci] if ci < len(row) else ""
                    cells[ci].width = widths[ci]
                    par = cells[ci].paragraphs[0]
                    par.paragraph_format.space_after = Pt(1)
                    par.paragraph_format.space_before = Pt(1)
                    add_rich(par, text, size=MONO_TABLE)
                    for r in par.runs:
                        r.bold = True if ri == 0 or text.startswith("**") else r.bold
                    if ri == 0:
                        shade(cells[ci], "EFE7F0")
                    elif text.startswith("**ИТОГО") or text.startswith("**2 530"):
                        shade(cells[ci], "F7F3F8")
                    # числовые колонки — по правому краю
                    if ci and re.search(r"\d[\d  ]*$", text.replace("**", "").replace("*", "")):
                        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            # шапка повторяется на новой странице
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            tr_pr.append(tr_pr.makeelement(qn("w:tblHeader"), {}))
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        if stripped.startswith("---"):
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            if level == 1:
                par = doc.add_paragraph()
                par.paragraph_format.space_after = Pt(6)
                add_rich(par, text, size=19, color=ACCENT)
                for r in par.runs:
                    r.bold = True
            else:
                par = doc.add_paragraph()
                par.paragraph_format.space_before = Pt(10)
                par.paragraph_format.space_after = Pt(4)
                add_rich(par, text, size=13, color=ACCENT if level == 2 else GREY)
                for r in par.runs:
                    r.bold = True
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            text = re.sub(r"^\d+\.\s", "", stripped)
            par = doc.add_paragraph(style="List Number")
            par.paragraph_format.space_after = Pt(2)
            add_rich(par, text)
            i += 1
            continue

        if stripped.startswith("- "):
            text = stripped[2:]
            par = doc.add_paragraph(style="List Bullet")
            par.paragraph_format.space_after = Pt(2)
            add_rich(par, text)
            i += 1
            continue

        par = doc.add_paragraph()
        par.paragraph_format.space_after = Pt(4)
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_rich(par, stripped, color=GREY if stripped.startswith("**") is False else None)
        i += 1

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_rich(
        footer,
        "SM Production · Москва, Остаповский проезд, д. 9, с. 13 · +7 495 799-00-19 · "
        "sm-production.ru · Расчёт от 22.10.2026, действует до 05.11.2026",
        size=8,
        color=GREY,
    )
    doc.save(str(out_path))
    print(f"OK: {out_path}")


if __name__ == "__main__":
    build(Path(sys.argv[1]), Path(sys.argv[2]))
