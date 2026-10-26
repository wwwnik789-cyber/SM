#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Проверка PPTX без PowerPoint: рисуем превью каждого слайда из самого файла
(позиции, картинки, тексты) — ловим наезды и переполнение.
Результат: /tmp/pptx-preview/pp-XX.png + контактный лист.
"""
import io, os, glob
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

BASE = os.path.dirname(os.path.abspath(__file__))
FONTS = os.path.join(BASE, "fonts")
PPTX = os.path.join(BASE, "SM-Production_Brending-avtosalonov.pptx")
OUTDIR = "/tmp/pptx-preview"
os.makedirs(OUTDIR, exist_ok=True)

PXE = 9525
W, H = 1280, 720

FONTFILE = {
    ("Unbounded", False): "Unbounded-400.ttf",
    ("Unbounded", True): "Unbounded-700.ttf",
    ("Unbounded ExtraBold", False): "Unbounded-800.ttf",
    ("Manrope", False): "Manrope-400.ttf",
    ("Manrope", True): "Manrope-700.ttf",
    ("Manrope SemiBold", False): "Manrope-600.ttf",
    ("Manrope ExtraBold", False): "Manrope-800.ttf",
}
_fc = {}
def font(name, bold, size_px):
    k = (name, bold, int(round(size_px)))
    if k not in _fc:
        f = FONTFILE.get((name, bool(bold)), "Manrope-400.ttf")
        _fc[k] = ImageFont.truetype(os.path.join(FONTS, f), max(1, int(round(size_px))))
    return _fc[k]

def emu2px(v):
    return v / PXE

def shape_rgb(shp):
    try:
        if shp.fill.type is not None and shp.fill.type == 1:
            c = shp.fill.fore_color.rgb
            return (c[0], c[1], c[2])
    except Exception:
        pass
    return None

def shape_alpha(shp):
    try:
        xml = shp.fill._xPr.xml
    except Exception:
        return 100
    import re
    m = re.search(r'<a:alpha val="(\d+)"', xml)
    return int(m.group(1)) / 1000 if m else 100

prs = Presentation(PPTX)
sx = emu2px(prs.slide_width) / W
sy = emu2px(prs.slide_height) / H

for idx, slide in enumerate(prs.slides, 1):
    canvas = Image.new("RGB", (W, H), (10, 10, 13))
    d = ImageDraw.Draw(canvas)
    for shp in slide.shapes:
        try:
            x = emu2px(shp.left) / sx; y = emu2px(shp.top) / sy
            w = emu2px(shp.width) / sx; h = emu2px(shp.height) / sy
        except TypeError:
            continue

        if shp.shape_type == 13 or shp.__class__.__name__ == "Picture":  # picture
            blob = shp.image.blob
            im = Image.open(io.BytesIO(blob)).convert("RGB").resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
            # скругление
            try:
                prst = shp._element.spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom').get("prst")
            except Exception:
                prst = "rect"
            if prst == "roundRect":
                r = int(min(w, h) * 0.06)
                mask = Image.new("L", im.size, 0)
                ImageDraw.Draw(mask).rounded_rectangle([0, 0, im.size[0] - 1, im.size[1] - 1], radius=r, fill=255)
                canvas.paste(im, (int(x), int(y)), mask)
            else:
                canvas.paste(im, (int(x), int(y)))
            continue

        # автофигура
        if shp.has_text_frame and not shp.text_frame.text.strip():
            rgb = shape_rgb(shp)
            if rgb:
                a = shape_alpha(shp)
                if a < 100:
                    layer = Image.new("RGB", (max(1, int(w)), max(1, int(h))), rgb)
                    mask = Image.new("L", layer.size, int(255 * a / 100))
                    canvas.paste(layer, (int(x), int(y)), mask)
                else:
                    rad = 6
                    try:
                        rad = int(min(w, h) * (shp.adjustments[0] if shp.adjustments else 0))
                    except Exception:
                        pass
                    d.rounded_rectangle([x, y, x + w, y + h], radius=min(rad, int(min(w, h) / 2)),
                                        fill=rgb, outline=None)
            continue

        # текст
        tf = shp.text_frame
        ty = y
        for p in tf.paragraphs:
            if not p.runs:
                ty += 4
                continue
            run = p.runs[0]
            size_px = run.font.size.pt / 0.75 if run.font.size else 14
            name = run.font.name or "Manrope"
            bold = bool(run.font.bold)
            f = font(name, bold, size_px)
            asc, desc = f.getmetrics()
            # межстрочный
            ls = p.line_spacing
            step = size_px * 1.2
            if ls is not None:
                try:
                    step = ls.pt / 0.75
                except Exception:
                    step = size_px * float(ls)
            text = "".join(r.text for r in p.runs)
            color = (244, 242, 236)
            try:
                c = run.font.color.rgb
                color = (c[0], c[1], c[2])
            except Exception:
                pass
            align = str(p.alignment) if p.alignment is not None else "LEFT"
            tx = x
            if "CENTER" in align:
                tx = x + (w - f.getlength(text)) / 2
            elif "RIGHT" in align:
                tx = x + w - f.getlength(text)
            d.text((tx, ty), text, font=f, fill=color)
            ty += step
    canvas.save(f"{OUTDIR}/pp-{idx:02d}.png")

# контактный лист
files = sorted(glob.glob(f"{OUTDIR}/pp-*.png"))
cols, cw, ch = 3, 470, 264
rows = (len(files) + cols - 1) // cols
sheet = Image.new("RGB", (cols * cw, rows * ch), (28, 28, 32))
for i, f in enumerate(files):
    im = Image.open(f).resize((cw - 8, ch - 8), Image.LANCZOS)
    sheet.paste(im, ((i % cols) * cw + 4, (i // cols) * ch + 4))
sheet.save("/tmp/pptx-preview/sheet.jpg", quality=88)
print("слайдов:", len(files), "->", OUTDIR)
