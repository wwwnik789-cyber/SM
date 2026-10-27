#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown → DOCX для документов BAK MASTERS (правится в Word/LibreOffice)."""
import os
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

INK = RGBColor(0x2B, 0x26, 0x21)
MUTED = RGBColor(0x7A, 0x6F, 0x63)
ACC = RGBColor(0x2F, 0x6F, 0x6B)
FONT = 'Manrope'


def set_cell_bg(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def clean(text):
    t = text
    t = re.sub(r'`([^`]+)`', r'\1', t)
    t = re.sub(r'\*\*([^*]+)\*\*', r'\1', t)
    return t.replace('⚠️', '⚠').strip()


def add_runs(par, text, size=10, bold=False, color=INK):
    """минимальная инлайн-разметка: **жирный**, `код`"""
    pos = 0
    for m in re.finditer(r'\*\*([^*]+)\*\*|`([^`]+)`', text):
        if m.start() > pos:
            r = par.add_run(text[pos:m.start()])
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.bold = bold
            r.font.name = FONT
        r = par.add_run(m.group(1) or m.group(2))
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.bold = True if m.group(1) else False
        r.font.name = 'DejaVu Sans Mono' if m.group(2) else FONT
        pos = m.end()
    if pos < len(text):
        r = par.add_run(text[pos:])
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.bold = bold
        r.font.name = FONT


def build(md_path, docx_path, title, subtitle='', meta=None):
    doc = Document()

    # поля страницы
    for s in doc.sections:
        s.left_margin = Cm(2.0)
        s.right_margin = Cm(2.0)
        s.top_margin = Cm(2.0)
        s.bottom_margin = Cm(1.8)

    st = doc.styles['Normal']
    st.font.name = FONT
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.3

    # ---- титульная часть (не отдельная страница) ----
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = INK
    r.font.name = FONT

    if subtitle:
        p = doc.add_paragraph()
        r = p.add_run(subtitle)
        r.font.size = Pt(11)
        r.font.color.rgb = MUTED
        r.font.name = FONT
    for m in (meta or []):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(m)
        r.font.size = Pt(9)
        r.font.color.rgb = MUTED
        r.font.name = FONT

    doc.add_paragraph()

    lines = open(md_path, encoding='utf-8').read().split('\n')
    i = 0
    first_h1_skipped = False
    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()
        if not s:
            i += 1
            continue

        # горизонтальная линия → пустой абзац с границей снизу
        if re.fullmatch(r'-{3,}', s):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '4')
            bottom.set(qn('w:color'), 'E4DBCD')
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # код-блок → серый блок с моношрифтом
        if s.startswith('```'):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            i += 1
            for b in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.4)
                pPr = p._p.get_or_add_pPr()
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:fill'), 'F2EEE7')
                pPr.append(shd)
                r = p.add_run(b.replace('\t', '    ') if b.strip() else ' ')
                r.font.name = 'DejaVu Sans Mono'
                r.font.size = Pt(9)
                r.font.color.rgb = INK
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # заголовки
        m = re.match(r'^(#{1,4})\s+(.*)$', s)
        if m:
            lvl = len(m.group(1))
            txt = clean(m.group(2))
            if lvl == 1 and not first_h1_skipped:
                first_h1_skipped = True
                i += 1
                continue            # H1 уже на титуле
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(txt)
            r.font.name = FONT
            r.font.bold = True
            r.font.color.rgb = ACC if (lvl <= 2 or txt.upper().startswith('БЛОК')) else INK
            r.font.size = Pt({1: 15, 2: 13.5, 3: 11.5, 4: 10.5}[lvl])
            i += 1
            continue

        # таблица
        if s.startswith('|') and s.endswith('|') and i + 1 < len(lines) and \
                re.fullmatch(r'\|[\s:|-]+\|', lines[i + 1].strip()):
            def cells(l):
                return [c.strip() for c in l.strip().strip('|').split('|')]
            rows = [cells(s)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(cells(lines[i]))
                i += 1
            ncol = len(rows[0])
            tbl = doc.add_table(rows=0, cols=ncol)
            tbl.style = 'Table Grid'
            tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
            for ri, row in enumerate(rows):
                row = (row + [''] * ncol)[:ncol]
                tr = tbl.add_row()
                for ci, val in enumerate(row):
                    cell = tr.cells[ci]
                    cell.text = ''
                    par = cell.paragraphs[0]
                    par.paragraph_format.space_after = Pt(2)
                    add_runs(par, val, size=9,
                             bold=(ri == 0),
                             color=MUTED if ri == 0 else INK)
                    if ri == 0:
                        set_cell_bg(cell, 'F6F1E9')
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # цитата
        if s.startswith('>'):
            buf = []
            while i < len(lines) and lines[i].strip().startswith('>'):
                buf.append(lines[i].strip().lstrip('>').strip())
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            left.set(qn('w:val'), 'single')
            left.set(qn('w:sz'), '18')
            left.set(qn('w:space'), '8')
            left.set(qn('w:color'), '2F6F6B')
            pbdr.append(left)
            pPr.append(pbdr)
            shd = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:fill'), 'EFEAE1')
            pPr.append(shd)
            add_runs(p, clean(' '.join(x for x in buf if x)), size=9, color=MUTED)
            continue

        # списки
        if re.match(r'^[-*]\s+', s):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, clean(re.sub(r'^[-*]\s+', '', s)), size=10)
            i += 1
            continue
        if re.match(r'^\d+\.\s+', s):
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, clean(re.sub(r'^\d+\.\s+', '', s)), size=10)
            i += 1
            continue

        # абзац
        buf = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith('#') or nxt.startswith('>')
                    or nxt.startswith('```') or (nxt.startswith('|') and nxt.endswith('|'))
                    or re.match(r'^[-*]\s+', nxt) or re.match(r'^\d+\.\s+', nxt)
                    or re.fullmatch(r'-{3,}', nxt)):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_runs(p, clean(' '.join(buf)), size=10)

    doc.save(docx_path)
    return docx_path


if __name__ == '__main__':
    print('модуль')
